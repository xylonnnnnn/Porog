"""Build the Russian academic report from docs/REPORT.md. Requires python-docx and Pillow.

Optional docs/assets/report-pages.json contains verified PDF page numbers for the static TOC.
Run the document skill renderer after generation; generation alone is not visual QA.
"""
from pathlib import Path
import json
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'docs/assets'
ASSETS.mkdir(parents=True, exist_ok=True)


def diagram():
    image = Image.new('RGB', (1700, 1040), 'white')
    draw = ImageDraw.Draw(image)
    font_path = '/System/Library/Fonts/Supplemental/Arial.ttf'
    if not Path(font_path).exists():
        font_path = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    normal = ImageFont.truetype(font_path, 29)
    small = ImageFont.truetype(font_path, 23)
    def box(bounds, title, detail):
        draw.rounded_rectangle(bounds, radius=18, fill='#F0F3F0', outline='#56675D', width=3)
        x1,y1,x2,y2 = bounds
        draw.text(((x1+x2)/2, y1+29), title, fill='#1F2923', font=normal, anchor='mt')
        draw.text(((x1+x2)/2, y1+81), detail, fill='#47574C', font=small, anchor='mt')
    def arrow(start,end,label=None):
        draw.line([start,end], fill='#56675D', width=4)
        x,y=end; sx,sy=start
        if y>sy: triangle=[(x,y),(x-10,y-18),(x+10,y-18)]
        elif x>sx: triangle=[(x,y),(x-18,y-10),(x-18,y+10)]
        else: triangle=[(x,y),(x+18,y-10),(x+18,y+10)]
        draw.polygon(triangle,fill='#56675D')
        if label: draw.text(((sx+x)/2+14,(sy+y)/2),label,font=small,fill='#47574C',anchor='lm')
    box((440,35,1260,175),'Porog.Desktop','Расписание · статус · трей · Отход')
    arrow((850,175),(850,300),'JSON через named pipe и ACL')
    box((440,300,1260,440),'Porog.Service','Планировщик · IPC · Controller')
    arrow((1260,370),(1400,370))
    box((1400,290,1680,450),'JSON','config · runtime')
    arrow((600,440),(600,585))
    arrow((1100,440),(1100,585))
    box((260,585,770,735),'HostsBlocker','Собственные маркеры hosts')
    box((890,585,1400,735),'ProcessBlocker','Проверка владельца и exe')
    arrow((515,735),(515,855))
    arrow((1145,735),(1145,855))
    box((260,855,770,995),'Домены','Явные имена IPv4 и IPv6')
    box((890,855,1400,995),'Процессы','Только SID владельца')
    image.save(ASSETS / 'architecture.png')


def set_font(style, size, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Times New Roman')
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if 'theme' in key.lower(): del fonts.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+key), 'Times New Roman')
    for border in style.element.xpath('.//w:pBdr'): border.getparent().remove(border)


def build():
    diagram()
    text = (ROOT / 'docs/REPORT.md').read_text()
    lines = text.splitlines()
    headings = [line[3:] for line in lines if line.startswith('## ')]
    pages_path = ASSETS / 'report-pages.json'
    pages = json.loads(pages_path.read_text()) if pages_path.exists() else {}
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = Cm(3); section.right_margin = Cm(1.5)
    section.header_distance = section.footer_distance = Cm(1)
    section.different_first_page_header_footer = True
    for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Caption']:
        set_font(doc.styles[name], 14, name.startswith('Heading'))
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.5; normal.space_after = Pt(0)
    normal.first_line_indent = Cm(1.25)
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.widow_control = True
    for name in ['Heading 1','Heading 2']:
        fmt = doc.styles[name].paragraph_format
        fmt.first_line_indent = Cm(0); fmt.space_before = Pt(16); fmt.space_after = Pt(10)
        fmt.line_spacing = 1.15; fmt.keep_with_next = True
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Subtitle'].font.italic = False
    for name in ['Title','Subtitle','Caption']:
        doc.styles[name].paragraph_format.first_line_indent = Cm(0)
    doc.styles['Caption'].font.size = Pt(11)
    doc.styles['Caption'].paragraph_format.line_spacing = 1.15
    doc.styles['Caption'].paragraph_format.space_before = Pt(6)
    doc.styles['Caption'].paragraph_format.space_after = Pt(8)
    # Neutral title page without invented author, faculty, supervisor or signatures.
    p = doc.add_paragraph('ПРОЕКТНАЯ РАБОТА')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(65); p.paragraph_format.space_after = Pt(75)
    p = doc.add_paragraph('Отчёт о реализации\nпроекта Порог', 'Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(30)
    p = doc.add_paragraph('Локальный блокировщик сайтов\nи приложений для Windows', 'Subtitle')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('Версия реализации 1.0.0\nНа основании двух исходных документов')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(60)
    p = doc.add_paragraph('19 октября 2026 года')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(75)
    doc.add_page_break()
    p = doc.add_paragraph('Содержание', 'Heading 1')
    p.paragraph_format.space_before = Pt(0)
    for i,title in enumerate(headings):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), f'section_{i}')
        r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = title; r.append(t); link.append(r); p._p.append(link)
        p.add_run('\t' + str(pages.get(title, '')))
    doc.add_page_break()
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    bookmark = table_number = 0
    i = next(i for i,line in enumerate(lines) if line.startswith('## '))
    while i < len(lines):
        line = lines[i].strip()
        if not line: i += 1; continue
        if line.startswith('## '):
            p = doc.add_paragraph(line[3:], 'Heading 1')
            if line[3:] == 'Заключение': p.paragraph_format.page_break_before = True
            start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(bookmark)); start.set(qn('w:name'), f'section_{bookmark}')
            end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(bookmark))
            p._p.insert(0,start); p._p.append(end); bookmark += 1
            i += 1; continue
        if line.startswith('!['):
            p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(ASSETS / 'architecture.png'), width=Cm(16.2))
            shape = p._p.xpath('.//wp:docPr')[0]; shape.set('descr', 'Диаграмма компонентов Порог с UI, службой, JSON, hosts и процессами')
            p = doc.add_paragraph('Рисунок 1 — Диаграмма компонентов приложения', 'Caption')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in cells): rows.append(cells)
                i += 1
            table_number += 1
            captions = ['Функции обязательного объёма','Основные сущности данных','Пользовательские экраны','Результаты и границы проверок']
            p = doc.add_paragraph(f'Таблица {table_number} — {captions[table_number-1]}', 'Caption')
            p.paragraph_format.keep_with_next = True
            table = doc.add_table(rows=1, cols=len(rows[0])); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
            widths = [4.1,6.2,6.2] if table_number in (1,2) else [3.3,6.7,6.5]
            for column,width in zip(table.columns,widths): column.width = Cm(width)
            borders = OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el = OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            table._tbl.tblPr.append(borders)
            for ri,row in enumerate(rows):
                tr = table.rows[0] if ri == 0 else table.add_row()
                props = tr._tr.get_or_add_trPr(); props.append(OxmlElement('w:cantSplit'))
                if ri == 0: props.append(OxmlElement('w:tblHeader'))
                for ci,value in enumerate(row):
                    cell = tr.cells[ci]; cell.width = Cm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = value
                    tcpr = cell._tc.get_or_add_tcPr()
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7E7E7' if ri == 0 else 'FFFFFF'); tcpr.append(shade)
                    margins = OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        el = OxmlElement(f'w:{side}'); el.set(qn('w:w'),'100'); el.set(qn('w:type'),'dxa'); margins.append(el)
                    tcpr.append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.line_spacing = 1.15
                        p.paragraph_format.space_after = Pt(2); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        for run in p.runs: run.font.size = Pt(11); run.bold = ri == 0
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
            continue
        # Body paragraphs and source entries; retain plain technical names without Markdown residue.
        p = doc.add_paragraph(line.replace('`',''))
        if re.match(r'^\d+\. ',line):
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(8)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs: run.font.size = Pt(12)
        i += 1
    doc.core_properties.title = 'Отчёт о реализации проекта Порог'
    doc.core_properties.subject = 'Реализация локального блокировщика для Windows'
    doc.core_properties.author = ''
    doc.core_properties.keywords = 'Порог, Windows, MVP, расписание, отчёт'
    target = ROOT / 'docs/Отчёт_Порог.docx'; doc.save(target); print(target)


if __name__ == '__main__':
    build()
