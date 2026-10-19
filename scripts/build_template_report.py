"""Fill the supplied user-guide template, preserving opaque DOCX parts."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / 'Шаблон_отчета_по_реализации_проекта.docx'
OUT = ROOT / 'docs/Отчёт_Порог_по_шаблону.docx'
QA = ROOT / 'docs/qa/template-guide'
QA.mkdir(parents=True, exist_ok=True)
source = REF.read_bytes()
with ZipFile(BytesIO(source)) as archive:
    baseline = {name: archive.read(name) for name in archive.namelist()}
inventory = {name: {'size': len(data), 'sha256': sha256(data).hexdigest()} for name, data in baseline.items()}
(QA / 'package-baseline.json').write_text(json.dumps(inventory, indent=2), encoding='utf-8')
(QA / 'artifact.md').write_text(f'''# Contract for the Porog user guide

Reference: {REF}
SHA-256: {sha256(source).hexdigest()}
Reference rendering: ../template-reference/page-1.png through page-3.png, all inspected.
Three pages, one portrait A4 section. Section geometry and sectPr are preserve-only:
11906 × 16838 twips; margins left 1701, right 850, top/bottom 1134;
header/footer distance 708; single column; no actual headers, footers, tables or fields.

Source roles: title p[1], introduction p[3], section labels p[4], p[10], p[20],
p[32], p[42], p[52], p[57], all in word/document.xml / w:document/w:body.
Calibri 12 pt, black bold; source mixed punctuation fonts normalized to Calibri.
Normal line rhythm 278/240 and 160 twips after. Left aligned.
Bullet role p[5] numId 5 ilvl 0. Nested step roles p[25], p[37], p[47]
use numIds 7, 8, 9 ilvl 1. Existing numbering definitions remain untouched.
No borders, rules, coloured blocks, headers, page numbers or ornamental furniture.

Editable slots: replace all 62 source body paragraphs with completed content in the
same ordered sections 0–6. Remove author-facing prompts and the participant placeholder.
Reuse cloned body, bullet and step paragraph patterns; expand slots as required for
installation, three functions, examples, troubleshooting and five FAQ answers.
The title uses existing Title style with explicit source Calibri 12 black bold formatting;
section labels use existing Heading1 with the same explicit source formatting.
Keep headings with following content. Clone plain paragraphs for figures/captions;
embed unmodified real Windows screenshots, proportionally fit within source text area.
Captions identify demonstration data. Add real external hyperlinks to repository and build.
Numbering is native, with independent existing IDs for each function.

Preserve-only parts: all baseline parts except word/document.xml,
word/_rels/document.xml.rels and [Content_Types].xml. Preserve all preexisting
relationships; only append image and hyperlink relationships. New media allowed.
Source styles, theme, numbering, settings, properties, font table and web settings
must be byte-identical. No fields, controls, comments, footnotes, bookmarks,
tables or images in the source; no unsupported control content requires rewriting.

Acceptance: sections 0–6 complete, no prompts, three function names/purposes/steps/
examples, first-screen picture and caption, three or more errors, five FAQ answers,
clickable repository and Windows build. Render reference/final diff and inspect every
final PNG. Compare package inventory and exact section properties; reference hash unchanged.
''', encoding='utf-8')

doc = Document(REF)
patterns = [deepcopy(p._p) for p in doc.paragraphs]
body = doc._element.body
sect = deepcopy(body.sectPr)
for child in list(body):
    body.remove(child)
body.append(sect)

def para(text='', pattern=10, *, style=None, keep=False):
    element = deepcopy(patterns[pattern])
    for child in list(element):
        if child.tag != qn('w:pPr'):
            element.remove(child)
    body.insert(len(body)-1, element)
    from docx.text.paragraph import Paragraph
    p = Paragraph(element, doc._body)
    if style:
        p.style = style
    f = p.paragraph_format
    f.space_before = Pt(0)
    f.space_after = Pt(8)
    f.line_spacing = 278 / 240
    f.keep_together = True
    f.keep_with_next = keep
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0,0,0)
    return p

def heading(text):
    return para(text, 3, style='Heading 1', keep=True)

def bullet(text):
    return para(text, 4)

def step(text, group):
    return para(text, {1:24, 2:36, 3:46}[group])

def link(label, url):
    p = para()
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
    for tag, attrs in [('rFonts', {'ascii':'Calibri','hAnsi':'Calibri'}), ('sz',{'val':'24'}), ('b',{}), ('color',{'val':'000000'}), ('u',{'val':'single'})]:
        el=OxmlElement('w:'+tag)
        for key,value in attrs.items(): el.set(qn('w:'+key),value)
        rp.append(el)
    r.append(rp); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r);p._p.append(h)
    return p

def image(name, caption, width=14.0):
    p=para('', keep=True)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(4)
    picture=p.add_run().add_picture(str(ROOT/'docs/evidence/windows/screenshots'/name),width=Cm(width))
    picture._inline.docPr.set('descr', caption)
    cp=para(caption)
    cp.paragraph_format.space_after=Pt(12)

para('Порог инструкция пользователя',0,style='Title',keep=True)
para('Пошаговое руководство по работе с MVP',1,keep=True)
para('Документ поможет установить «Порог», настроить первый рабочий промежуток, получить блокировку отвлекающих ресурсов и запустить короткий перерыв «Отход». Выполняйте разделы по порядку; исходный код и готовая Windows-сборка доступны по ссылкам ниже.',2)

heading('0. ВВЕДЕНИЕ')
bullet('Название проекта: «Порог» (Porog).')
bullet('Для кого этот продукт: для студентов и других пользователей личного компьютера Windows, которым нужно заранее выделять время на учёбу или работу и ограничивать отвлечения.')
bullet('Что он делает: по расписанию ограничивает выбранные сайты и приложения, показывает состояние рабочего блока и предлагает полноэкранный перерыв от 3 до 12 минут.')
bullet('Сколько времени займёт знакомство: ориентировочно 10–15 минут на установку и настройку, в зависимости от скачивания; минимальный пробный «Отход» длится ещё 3 минуты.')
bullet('Платформа: Windows 10/11 x64; права администратора нужны для установки службы. Проект предназначен для одного владельца на личном компьютере.')
para('Репозиторий проекта на GitHub:',keep=True)
link('https://github.com/xylonnnnnn/Porog','https://github.com/xylonnnnnn/Porog')
para('Ссылка на работающий MVP — готовая Windows-сборка:',keep=True)
link('Скачать Windows x64 из артефакта GitHub Actions','https://github.com/xylonnnnnn/Porog/actions/runs/39228874422/artifacts/12176386932')
para('Репозиторий закрытый: для скачивания войдите в GitHub под учётной записью с доступом. Если артефакт уже недоступен, в репозитории откройте Actions → Build and verify Porog, запустите workflow через Run workflow и скачайте новый артефакт Porog-win-x64 после успешной сборки.')

heading('1. ПЕРВЫЙ ВХОД В ПРОДУКТ')
para('Как попасть в продукт и выполнить установку:',keep=True)
bullet('Скачайте Windows-сборку по ссылке выше. Распакуйте скачанный архив; если внутри находится Porog-win-x64.zip, распакуйте и его. В итоговой папке должны быть каталоги Service, Desktop и scripts. Не запускайте программу прямо из архива.')
bullet('Откройте меню «Пуск», найдите Windows PowerShell и выберите «Запуск от имени администратора». Перейдите в распакованную папку командой cd. Например, для папки C:\\Porog выполните:')
para("cd 'C:\\Porog'")
para('.\\scripts\\Install.ps1')
bullet('Дождитесь сообщения Installed. Установщик создаст службу Porog, автозапуск интерфейса и ярлык в меню «Пуск». Если главное окно не появилось, откройте Porog через «Пуск» или дважды нажмите значок в трее рядом с часами.')
para('При установке используйте повышение прав своей учётной записи. Если Windows требует пароль другого администратора, сначала получите SID владельца командой whoami /user в обычном PowerShell. Администратор должен передать полученный SID параметром -OwnerSid установщику.')
para('Что нужно на старте: регистрация в самом приложении не требуется; личные данные вводить не нужно. На первом запуске проверьте состояние службы. Если расписание ещё не создано, добавьте промежуток по разделу 2.')
para('Что пользователь увидит: экран «Обзор», состояние службы, сведения о текущем или ближайшем промежутке и боковое меню «Расписание», «Отход», «Диагностика». На рисунках показаны данные тестового примера Diploma focus; при вводе выбирайте будущие даты и время на своём компьютере.',keep=True)
image('01-status.png','Рисунок 1. Стартовый раздел «Обзор» с примером уже сохранённого будущего промежутка.')

heading('2. ОСНОВНАЯ ФУНКЦИЯ №1 (Core-функция)')
bullet('Название функции: блокировка сайтов и приложений на время работы.')
bullet('Зачем она нужна: помогает соблюдать заранее принятое решение о работе, ограничивая выбранные отвлекающие ресурсы до конца промежутка.')
para('Пошаговая инструкция:',keep=True)
step('Откройте «Расписание» → «Новый промежуток». Введите название, например «Подготовка к экзамену», и необязательное намерение «Решить три задачи». Установите будущее начало и более позднее окончание.',1)
step('В поле сайтов введите по одному домену на строку. В поле приложений укажите имена файлов .exe. Оставьте флажок «Промежуток включён». Сохраните работу в выбранных приложениях: их процессы будут принудительно завершаться.',1)
image('03-editor.png','Рисунок 2. Шаги 1–2: форма нового промежутка до заполнения названия, сайтов и приложений.',12.4)
bullet('Пример входных данных: название «Подготовка к экзамену»; начало завтра в 09:30; окончание завтра в 10:15; сайты youtube.com, www.youtube.com, m.youtube.com — каждый с новой строки; приложение Telegram.exe.')
para('Указывайте домен без https://, пути, порта и символа *. Каждый поддомен добавляется отдельно. Для приложения нужен файл Telegram.exe, а не путь к папке. Системные процессы, компоненты Porog и браузеры целиком выбрать нельзя.')
step('Нажмите «Сохранить» и убедитесь, что в «Расписании» появилась карточка с нужными датами, количеством доменов и процессов. После успешного сохранения данные применяет служба; отдельной кнопки запуска блока нет.',1)
image('02-schedules.png','Рисунок 3. Шаг 3: сохранённый пример содержит три домена и один процесс.')
step('К началу промежутка оставьте компьютер включённым, а службу — работающей. Перейдите в «Обзор» (расположение статуса — на рисунке 1) и дождитесь состояния «Закрыто до HH:MM». Проверьте выбранный сайт и повторный запуск указанного приложения.',1)
para('Пример результата: с 09:30 до 10:15 применены ограничения для трёх указанных доменов; Telegram.exe завершается при обнаружении, в том числе после повторного запуска. «Обзор» показывает срок окончания и оставшееся время. В 10:15 ограничения снимаются автоматически.')
para('Во время активного промежутка нельзя изменить, удалить, выключить или сократить его через приложение. Планируйте ограничения до старта. Доступность сайта может зависеть от DNS, кэша браузера и уже открытых соединений; порядок проверки приведён в разделе 5.')

heading('3. ОСНОВНАЯ ФУНКЦИЯ №2')
bullet('Название функции: планирование нескольких рабочих промежутков.')
bullet('Зачем она нужна: позволяет заранее распределить время на разные задачи и подготовить ограничения, например вечером на следующий день.')
para('Пошаговая инструкция:',keep=True)
step('Откройте «Расписание». Нажмите «Новый промежуток» для следующей задачи либо «Изменить» на карточке будущего промежутка. Список и кнопки показаны на рисунке 3.',2)
step('В форме задайте название, намерение, обе даты и время начала и окончания, затем списки сайтов и приложений. Время относится к часовому поясу Windows, указанному в форме. Для работы после полуночи выберите следующий день в дате окончания.',2)
image('03-editor.png','Рисунок 4. Шаг 2: дата и время начала и окончания задаются отдельно; в тестовой системе часовой пояс UTC.',10.7)
step('Нажмите «Сохранить». Вернитесь в «Расписание» и проверьте карточки; затем откройте «Обзор», чтобы увидеть ближайший старт, как на рисунке 1. Повторите добавление для других задач. Пока промежуток не начался, его можно изменить или удалить.',2)
bullet('Пример входных данных: завтра 09:30–10:15 — «Подготовка к экзамену»; завтра 10:30–11:15 — «Практическое задание». Намерение для второго: «Завершить первый раздел». Списки ограничений задаются для каждого промежутка.')
bullet('Пример результата: две сохранённые карточки; в «Обзоре» ближайший старт 09:30. За минуту до каждого старта приложение отправляет уведомление Windows, если интерфейс запущен и ОС разрешает уведомления.')
para('Включённые промежутки не должны пересекаться. Вариант 10:00–10:45 конфликтует с 09:30–10:15; измените начало на 10:15 или позже. Соседние интервалы разрешены. Если компьютер включён в середине запланированного блока, служба применит его на оставшееся время; полностью пропущенный блок задним числом не запускается.')

heading('4. ОСНОВНАЯ ФУНКЦИЯ №3')
bullet('Название функции: полноэкранный перерыв «Отход».')
bullet('Зачем она нужна: даёт повод отойти от экрана после работы, пройтись или выпить воды и завершает рабочий цикл коротким перерывом.')
para('Пошаговая инструкция:',keep=True)
step('Сохраните текущую работу и откройте раздел «Отход». Выберите целое число минут от 3 до 12. Для первого знакомства установите 3 минуты; на снимке показан пример с 5 минутами.',3)
image('04-break-settings.png','Рисунок 5. Шаг 1: выбор продолжительности «Отхода» перед запуском.')
step('Нажмите «Начать отход». На подключённых мониторах появится полноэкранное окно с таймером и предложением встать, пройтись и выпить воды. Кнопки пропуска нет.',3)
image('06-break.png','Рисунок 6. Шаг 2: запущенный «Отход» с обратным отсчётом; пример на 5 минут.')
step('Отойдите от компьютера и дождитесь нуля. По окончании полноэкранные окна закрываются автоматически, после чего можно продолжить работу.',3)
bullet('Пример входных данных: длительность 5 минут и нажатие «Начать отход».')
bullet('Пример результата: таймер начинает отсчёт с 05:00, во время перерыва показывает оставшееся время, по истечении пяти минут экран снова доступен. На рисунке 6 показан момент 04:59.')
para('Обычное закрытие окна во время «Отхода» блокируется. При временной потере связи со службой таймер идёт до ранее полученного срока. Если интерфейс аварийно завершился, откройте его снова: незавершённый «Отход» восстановится. Системные механизмы Windows и действия администратора приложение не блокирует.')

heading('5. ЧТО ДЕЛАТЬ, ЕСЛИ ЧТО-ТО ПОШЛО НЕ ТАК')
para('Начните с раздела «Диагностика»: проверьте связь со службой, доступность hosts, контроля процессов и «Отхода». После устранения причины используйте повторное подключение.',keep=True)
image('05-diagnostics.png','Рисунок 7. «Диагностика»: пример доступной службы и компонентов блокировки.')
bullet('Ошибка №1: «Служба недоступна». Что делать: нажмите повторное подключение в «Диагностике». В PowerShell проверьте Get-Service Porog. Если служба остановлена, администратор может выполнить Start-Service Porog. Если службы нет, выполните установку из раздела 1; при разных учётных записях проверьте SID владельца.')
bullet('Ошибка №2: промежуток не сохраняется — SCHEDULE_CONFLICT или сообщение о времени. Что делать: задайте будущее начало, окончание позже начала и устраните пересечение с другим включённым промежутком. Для работы через полночь измените дату окончания. После CONFIG_VERSION_CONFLICT откройте список заново и повторите редактирование.')
bullet('Ошибка №3: сайт доступен или приложение не закрывается. Что делать: убедитесь, что блок активен, а при настройке указаны точные домены и имена .exe. Закройте старые вкладки и проверьте защищённый DNS: он может обходить hosts. Списки исправляются до старта или после окончания блока. При HOSTS_WRITE_FAILED или PROCESS_CONTROL_FAILED сообщите администратору код из «Диагностики».')
bullet('Ошибка №4: PowerShell запрещает Install.ps1. Что делать: проверьте Get-ExecutionPolicy -List. На личном ПК для текущего окна выполните Set-ExecutionPolicy -Scope Process RemoteSigned. Скачанный сценарий просмотрите и при доверии разблокируйте в свойствах файла. Затем повторите установку. Ограничения групповой политики снимает администратор.')
bullet('Ошибка №5: ACTIVE_SESSION_LOCKED при изменении или удалении. Что делать: дождитесь конца рабочего промежутка. Это предусмотренное ограничение; редактирование доступно до старта. При CONFIG_CORRUPTED сохраните сообщение и обратитесь к администратору: порядок восстановления из резервной копии описан в README и руководстве репозитория.')

heading('6. ЧАСТЫЕ ВОПРОСЫ (FAQ)')
bullet('Можно ли отменить рабочий блок досрочно? Через интерфейс нельзя: он действует до окончания. До старта ошибочный промежуток можно изменить или удалить.')
bullet('Что будет после закрытия окна или перезагрузки? Крестик скрывает окно в трей; двойной щелчок возвращает его. Служба продолжает работать, а после включения заново определяет состояние расписания. Для уведомлений и «Отхода» нужен работающий интерфейс.')
bullet('Нужны ли аккаунт и интернет? После скачивания Porog работает локально без аккаунта. Данные и резервные копии хранятся в %ProgramData%\\Porog; история посещений не записывается. GitHub нужен только для получения сборки из закрытого репозитория.')
bullet('Почему нет предупреждения за минуту? Проверьте, запущен ли интерфейс, разрешены ли уведомления Porog в Windows и выключен ли режим «Не беспокоить». Отсутствие уведомления не отменяет блокировку.')
bullet('Как удалить Порог? После блока и «Отхода» выполните .\\scripts\\Uninstall.ps1 из папки сборки в PowerShell администратора. Данные сохранятся.')

buffer=BytesIO();doc.save(buffer)
with ZipFile(BytesIO(buffer.getvalue())) as generated:
    final={name:generated.read(name) for name in generated.namelist()}
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
for name,data in baseline.items():
    if name not in editable:
        final[name]=data
with ZipFile(OUT,'w',ZIP_DEFLATED) as archive:
    for name,data in final.items():archive.writestr(name,data)

with ZipFile(OUT) as archive:
    assert all(archive.read(name)==data for name,data in baseline.items() if name not in editable)
    check=Document(OUT)
    assert len(check.sections)==1
    from lxml import etree
    assert etree.tostring(check._element.body.sectPr)==etree.tostring(sect)
    assert len(check.inline_shapes)==7
assert sha256(REF.read_bytes()).hexdigest()==sha256(source).hexdigest()
print(OUT)
print('Preserve-only package parts and source section unchanged; 7 actual screenshots embedded.')
